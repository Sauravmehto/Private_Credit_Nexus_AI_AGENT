from __future__ import annotations

import csv
import io
import json as _json
import re
from datetime import date, datetime, timedelta, timezone
from typing import Any

from fastapi import HTTPException, UploadFile

from api.schemas import FileUploadResponse, PortfolioUploadResponse

# ---------------------------------------------------------------------------
# Schema validation
# ---------------------------------------------------------------------------

# Required columns and accepted aliases (all lowercased, stripped of punctuation
# for fuzzy matching). Key = canonical display name used in error messages.
REQUIRED_COLUMNS: dict[str, list[str]] = {
    "Borrower / Company": [
        "borrower", "company", "name", "borrower name", "deal name",
        "company name", "fund name", "entity", "issuer",
    ],
    "Instrument Type": [
        "instrument type", "instrument", "type", "product",
        "debt type", "facility type",
    ],
    "Principal ($)": [
        "principal", "principal ($)", "loan amount", "commitment",
        "deal size", "facility size", "amount", "loan size",
        "face value", "investment amount",
    ],
    "Interest Rate": [
        "interest rate", "rate", "coupon", "cash coupon",
        "allin rate", "yield", "rate (%)", "annual rate",
    ],
    "Tenure (Mo.)": [
        "tenure", "tenor", "term", "tenure (mo.)", "tenor (mo.)",
        "term (mo.)", "term months", "duration", "loan term",
    ],
    "Orig. Fee": [
        "orig. fee", "origination fee", "orig fee", "upfront fee",
        "arrangement fee", "front end fee",
    ],
    "EOT Fee": [
        "eot fee", "end of term fee", "exit fee", "eot. fee",
        "maturity fee", "back end fee",
    ],
    "Covenant Status": [
        "covenant status", "status", "health", "credit status",
        "covenant", "compliance",
    ],
}


def _norm(s: str) -> str:
    """Lowercase and strip non-alphanumeric characters."""
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _check_schema(columns: list[str]) -> list[str]:
    """Return list of canonical display names for any missing required columns."""
    normed_cols = [_norm(c) for c in columns if c]
    missing: list[str] = []
    for display_name, aliases in REQUIRED_COLUMNS.items():
        normed_aliases = [_norm(a) for a in aliases]
        matched = any(
            alias == col
            or (len(alias) >= 4 and col and alias in col)
            or (len(col) >= 4 and alias and col in alias)
            for alias in normed_aliases
            for col in normed_cols
        )
        if not matched:
            missing.append(display_name)
    return missing


def _extract_xlsx_headers(content: bytes) -> list[str]:
    """
    Find the most likely header row in an xlsx workbook.
    Scans the first 20 rows looking for a row that contains known keywords.
    Returns the header strings found.
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is not installed. Run: pip install openpyxl",
        ) from exc

    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    if ws is None:
        return []

    keywords = {"borrower", "company", "name", "principal", "loan", "instrument", "amount", "deal", "rate", "tenure"}

    for row in ws.iter_rows(max_row=20, values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row]
        nonempty = [c for c in cells if c]
        if len(nonempty) < 2:
            continue
        row_text = " ".join(cells).lower()
        if any(kw in row_text for kw in keywords):
            return nonempty
        if len(nonempty) >= 4:
            return nonempty

    wb.close()
    return []


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

async def parse_upload(file: UploadFile) -> FileUploadResponse:
    filename = file.filename or ""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext not in {"csv", "xlsx", "xls", "xlsm", "json"}:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Unsupported file type. Please upload .xlsx, .csv, or .json.",
                "missing_columns": [],
            },
        )

    content = await file.read()

    # ── CSV ────────────────────────────────────────────────────────────────
    if ext == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        rows = [dict(row) for row in reader]
        columns = list(reader.fieldnames or [])

        missing = _check_schema(columns)
        if missing:
            raise HTTPException(
                status_code=422,
                detail={
                    "message": "Invalid file structure detected. Please upload the correct Excel, CSV, or JSON file.",
                    "missing_columns": missing,
                    "found_columns": columns,
                },
            )

        return FileUploadResponse(
            filename=filename,
            row_count=len(rows),
            columns=columns,
            sheet_name=None,
            preview=rows[:5],
        )

    # ── JSON ───────────────────────────────────────────────────────────────
    if ext == "json":
        import json as _json
        try:
            parsed = _json.loads(content)
        except _json.JSONDecodeError as exc:
            raise HTTPException(status_code=422, detail={"message": f"Invalid JSON: {exc}", "missing_columns": []}) from exc

        records: list[Any] = []
        if isinstance(parsed, list):
            records = parsed
        elif isinstance(parsed, dict):
            for key in ("deals", "portfolio", "data"):
                if isinstance(parsed.get(key), list):
                    records = parsed[key]
                    break

        if not records:
            raise HTTPException(status_code=422, detail={"message": "No data rows found in JSON.", "missing_columns": []})

        columns = list(records[0].keys()) if isinstance(records[0], dict) else []
        missing = _check_schema(columns)
        if missing:
            raise HTTPException(
                status_code=422,
                detail={
                    "message": "Invalid file structure detected. Please upload the correct Excel, CSV, or JSON file.",
                    "missing_columns": missing,
                    "found_columns": columns,
                },
            )

        return FileUploadResponse(
            filename=filename,
            row_count=len(records),
            columns=columns,
            sheet_name=None,
            preview=records[:5],
        )

    # ── XLSX / XLS ─────────────────────────────────────────────────────────
    try:
        import openpyxl
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is not installed. Run: pip install openpyxl",
        ) from exc

    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    # Score each sheet and pick the one richest in portfolio keywords
    _portfolio_kw2 = {"borrower", "company", "name", "principal", "instrument",
                      "amount", "rate", "tenure", "status", "loan", "deal"}
    _best_sn = wb.active.title if wb.active else wb.sheetnames[0]
    _best_sc = -1
    for _sn in wb.sheetnames:
        _ws_tmp = wb[_sn]
        _sc = 0
        for _r in _ws_tmp.iter_rows(max_row=20, values_only=True):
            _rt = " ".join(str(c).lower() for c in _r if c is not None)
            _sc += sum(1 for kw in _portfolio_kw2 if kw in _rt)
        if _sc > _best_sc:
            _best_sc = _sc
            _best_sn = _sn
    ws = wb[_best_sn]
    sheet_name = ws.title if ws else None

    # Find the real header row (first row with known keywords)
    header_row_values: tuple[Any, ...] | None = None
    keywords = {"borrower", "company", "name", "principal", "instrument", "amount", "rate", "tenure", "status"}
    rows_seen: list[tuple[Any, ...]] = []

    if ws:
        for row in ws.iter_rows(max_row=20, values_only=True):
            rows_seen.append(row)
            cells_str = [str(c).strip() if c is not None else "" for c in row]
            nonempty = [c for c in cells_str if c]
            if len(nonempty) < 2:
                continue
            row_text = " ".join(cells_str).lower()
            if any(kw in row_text for kw in keywords):
                header_row_values = row
                break
            if len(nonempty) >= 4 and header_row_values is None:
                header_row_values = row  # fallback: first wide row

    if header_row_values is None:
        raise HTTPException(status_code=422, detail={"message": "The uploaded file appears to be empty or has no recognisable header row.", "missing_columns": []})

    columns = [
        str(h).strip() if h is not None else f"Column_{i + 1}"
        for i, h in enumerate(header_row_values)
        if h is not None and str(h).strip()
    ]

    missing = _check_schema(columns)
    if missing:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Invalid file structure detected. Please upload the correct Excel, CSV, or JSON file.",
                "missing_columns": missing,
                "found_columns": columns,
            },
        )

    # Full data extraction
    header_idx = rows_seen.index(header_row_values)
    all_columns = [
        str(h).strip() if h is not None else f"Column_{i + 1}"
        for i, h in enumerate(header_row_values)
    ]
    data_rows: list[dict[str, Any]] = []
    if ws:
        for i, row in enumerate(ws.iter_rows(min_row=header_idx + 2, values_only=True)):
            nonempty = [v for v in row if v is not None and str(v).strip()]
            if not nonempty:
                continue
            data_rows.append({
                col: (val if isinstance(val, (int, float, bool)) else str(val) if val is not None else None)
                for col, val in zip(all_columns, row)
            })

    return FileUploadResponse(
        filename=filename,
        row_count=len(data_rows),
        columns=columns,
        sheet_name=sheet_name,
        preview=data_rows[:5],
    )


# ---------------------------------------------------------------------------
# Full parse — returns complete Deal[] shaped for the frontend
# ---------------------------------------------------------------------------

def _nk(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def _col(row: dict[str, Any], keys: list[str]) -> Any:
    normed: dict[str, Any] = {_nk(k): v for k, v in row.items()}
    for key in keys:
        nk = _nk(key)
        val = normed.get(nk)
        if val not in (None, "", "None"):
            return val
        for ck, cv in normed.items():
            if len(ck) >= 3 and len(nk) >= 3 and (ck in nk or nk in ck) and cv not in (None, "", "None"):
                return cv
    return ""


def _num(row: dict[str, Any], keys: list[str]) -> float:
    v = _col(row, keys)
    if v in ("", None):
        return 0.0
    try:
        return float(re.sub(r"[$,%\s,]", "", str(v)).replace("(", "-").replace(")", "")) or 0.0
    except ValueError:
        return 0.0


def _pct(row: dict[str, Any], keys: list[str]) -> float:
    v = _num(row, keys)
    return v / 100.0 if v > 1.0 else v


def _parse_date_val(orig_raw: Any, tenure: int) -> str:
    """Parse various date formats to ISO YYYY-MM-DD string."""
    if orig_raw not in (None, "", "None"):
        if isinstance(orig_raw, datetime):
            return orig_raw.strftime("%Y-%m-%d")
        if isinstance(orig_raw, date):
            return orig_raw.isoformat()
        # Excel serial number
        if isinstance(orig_raw, (int, float)) and float(orig_raw) > 40000:
            epoch = datetime(1899, 12, 30, tzinfo=timezone.utc)
            return (epoch + timedelta(days=int(orig_raw))).strftime("%Y-%m-%d")
        if isinstance(orig_raw, str) and orig_raw.strip():
            s = orig_raw.strip()
            for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y", "%m-%d-%Y"):
                try:
                    return datetime.strptime(s[:10], fmt).strftime("%Y-%m-%d")
                except ValueError:
                    continue
            try:
                return datetime.fromisoformat(s[:19]).strftime("%Y-%m-%d")
            except ValueError:
                pass
    months_ago = round(tenure * 0.4)
    return (datetime.now(timezone.utc) - timedelta(days=months_ago * 30)).strftime("%Y-%m-%d")


def _parse_repayment_structure(text: Any, tenure: int) -> tuple[int, int] | None:
    """Extract IO/amort months from phrases like 'IO 12 mo -> 24 mo amortizing'."""
    if text in (None, "", "None"):
        return None

    s = str(text).lower()
    io_match = re.search(r"\bio\s*(\d+)\s*(?:mo|month|months)?", s)
    amort_match = re.search(r"(\d+)\s*(?:mo|month|months)?\s*amort", s)

    io = int(io_match.group(1)) if io_match else None
    amort = int(amort_match.group(1)) if amort_match else None

    if io is not None and amort is not None:
        return (io, amort)
    if "bullet" in s:
        return (io or 0, 0)
    if io is not None:
        return (io, max(0, tenure - io))
    if amort is not None:
        return (max(0, tenure - amort), amort)
    return None


def _norm_row(row: dict[str, Any], idx: int) -> dict[str, Any]:
    """Normalize a raw Excel/CSV row to the frontend Deal interface shape."""
    name = str(
        _col(row, ["Borrower", "Company", "Name", "Borrower Name", "Deal Name",
                   "Company Name", "Fund Name", "Entity", "Issuer"]) or f"Deal {idx + 1}"
    )
    inst = str(
        _col(row, ["Instrument Type", "Instrument", "Type", "Product",
                   "Debt Type", "Facility Type"]) or "term loan"
    ).lower().strip()

    principal = _num(row, ["Principal", "Principal ($)", "Loan Amount", "Commitment",
                            "Deal Size", "Facility Size", "Amount", "Loan Size",
                            "Face Value", "Investment Amount"])
    if not principal:
        for v in row.values():
            try:
                fv = float(str(v).replace(",", "").replace("$", "").strip())
                if fv >= 10_000:
                    principal = fv
                    break
            except (ValueError, AttributeError):
                pass

    rate = _pct(row, ["Interest Rate", "Rate", "Coupon", "Cash Coupon",
                      "All-in Rate", "Yield", "Rate (%)", "Annual Rate"])
    if not rate:
        rate = 0.12

    tenure_raw = _col(row, ["Tenure", "Tenor", "Term", "Tenure (Mo.)", "Tenor (Mo.)",
                             "Term (Mo.)", "Term Months", "Duration", "Loan Term"])
    tenure = int(re.sub(r"[^0-9]", "", str(tenure_raw or "36")) or "36") or 36

    status_raw = str(
        _col(row, ["Covenant Status", "Status", "Health", "Covenant", "Credit Status"]) or "compliant"
    ).lower()
    if any(x in status_raw for x in ("breach", "default", "violation")):
        status = "breach risk"
    elif any(x in status_raw for x in ("monitor", "watch", "review")):
        status = "monitoring"
    else:
        status = "compliant"

    repayment = _parse_repayment_structure(
        _col(row, ["Repayment Structure", "Repayment", "Payment Structure", "Amortization Structure"]),
        tenure,
    )
    io_raw = _col(row, ["IO Period", "IO Months", "IO (Mo.)", "IO Period (Mo.)", "Interest Only", "IO"])
    io_str = re.sub(r"[^0-9]", "", str(io_raw or ""))
    if repayment:
        io, amort = repayment
    else:
        io = int(io_str) if io_str else max(0, round(tenure * 0.2))
        amort = max(0, tenure - io)

    orig_fee = _pct(row, ["Orig Fee", "Origination Fee", "Orig. Fee",
                           "Upfront Fee", "Arrangement Fee", "Front End Fee"])
    eot_fee = _pct(row, ["EOT Fee", "End of Term Fee", "Exit Fee",
                          "EoT Fee", "Maturity Fee", "Back End Fee"])
    pik_rate = _pct(row, ["PIK Rate", "PIK", "PIK (%)", "Payment-in-Kind"])
    if not orig_fee:
        orig_fee = 0.02
    if not eot_fee:
        eot_fee = 0.025

    orig_raw = _col(row, ["Origination Date", "Close Date", "Closing Date", "Issue Date",
                           "Start Date", "Funding Date", "Date", "Trade Date"])
    origination = _parse_date_val(orig_raw, tenure)

    outstanding = _num(row, ["Outstanding", "Outstanding ($)", "Outstanding Balance"]) or principal
    deal_id = str(_col(row, ["Deal ID", "ID", "Ref"]) or f"D{str(idx + 1).zfill(3)}")

    return {
        "id": f"D{str(idx + 1).zfill(3)}",
        "deal_id": deal_id,
        "name": name or f"Deal {idx + 1}",
        "instrument_type": inst,
        "principal": principal or 0.0,
        "outstanding": outstanding,
        "rate": rate,
        "io": io,
        "amort": amort,
        "tenure": tenure,
        "orig_fee": orig_fee,
        "eot_fee": eot_fee,
        "pik_rate": pik_rate,
        "warrant": str(_col(row, ["Warrant", "Warrant Coverage", "Warrants"]) or ""),
        "status": status,
        "sector": str(_col(row, ["Sector", "Industry", "Vertical"]) or ""),
        "stage": str(_col(row, ["Stage", "Company Stage", "Round"]) or ""),
        "maturity": str(_col(row, ["Maturity Date", "Maturity", "Due Date"]) or ""),
        "origination": origination,
        "interest_type": str(_col(row, ["Interest Type", "Rate Type"]) or "Fixed"),
    }


async def parse_upload_full(file: UploadFile) -> PortfolioUploadResponse:
    """Parse the uploaded file and return the complete normalized deals array."""
    filename = file.filename or ""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext not in {"csv", "xlsx", "xls", "xlsm", "json"}:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Unsupported file type. Please upload .xlsx, .csv, or .json.",
                "missing_columns": [],
            },
        )

    content = await file.read()
    raw_rows: list[dict[str, Any]] = []

    # ── CSV ────────────────────────────────────────────────────────────────
    if ext == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        raw_rows = [dict(row) for row in reader]
        columns = list(reader.fieldnames or [])
        missing = _check_schema(columns)
        if missing:
            raise HTTPException(
                status_code=422,
                detail={"message": "Invalid file structure.", "missing_columns": missing, "found_columns": columns},
            )

    # ── JSON ───────────────────────────────────────────────────────────────
    elif ext == "json":
        try:
            parsed = _json.loads(content)
        except _json.JSONDecodeError as exc:
            raise HTTPException(status_code=422, detail={"message": f"Invalid JSON: {exc}", "missing_columns": []}) from exc
        if isinstance(parsed, list):
            raw_rows = parsed
        elif isinstance(parsed, dict):
            for key in ("deals", "portfolio", "data"):
                if isinstance(parsed.get(key), list):
                    raw_rows = parsed[key]
                    break
        if not raw_rows:
            raise HTTPException(status_code=422, detail={"message": "No data rows found in JSON.", "missing_columns": []})
        columns = list(raw_rows[0].keys()) if isinstance(raw_rows[0], dict) else []
        missing = _check_schema(columns)
        if missing:
            raise HTTPException(
                status_code=422,
                detail={"message": "Invalid file structure.", "missing_columns": missing, "found_columns": columns},
            )

    # ── XLSX / XLS ─────────────────────────────────────────────────────────
    else:
        try:
            import openpyxl
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="openpyxl is not installed.") from exc

        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        if not wb.sheetnames:
            raise HTTPException(status_code=422, detail={"message": "Empty workbook.", "missing_columns": []})

        # Score each sheet by portfolio keyword density — pick the best sheet
        _portfolio_kw = {"borrower", "company", "name", "principal", "instrument",
                         "amount", "rate", "tenure", "status", "loan", "deal"}
        best_sheet_name = wb.active.title if wb.active else wb.sheetnames[0]
        best_score = -1
        for sn in wb.sheetnames:
            ws_tmp = wb[sn]
            score = 0
            for _row in ws_tmp.iter_rows(max_row=20, values_only=True):
                row_text = " ".join(str(c).lower() for c in _row if c is not None)
                score += sum(1 for kw in _portfolio_kw if kw in row_text)
            if score > best_score:
                best_score = score
                best_sheet_name = sn

        ws = wb[best_sheet_name]
        repayment_by_deal_id: dict[str, Any] = {}

        if "Deal Details" in wb.sheetnames:
            detail_ws = wb["Deal Details"]
            detail_header: tuple[Any, ...] | None = None
            detail_header_idx = 0
            for i, detail_row in enumerate(detail_ws.iter_rows(max_row=10, values_only=True), start=1):
                norm_cells = [_norm(str(c)) if c is not None else "" for c in detail_row]
                if "dealid" in norm_cells and "repaymentstructure" in norm_cells:
                    detail_header = detail_row
                    detail_header_idx = i
                    break

            if detail_header:
                norm_headers = [_norm(str(h)) if h is not None else "" for h in detail_header]
                deal_id_idx = norm_headers.index("dealid")
                repayment_idx = norm_headers.index("repaymentstructure")
                for detail_row in detail_ws.iter_rows(min_row=detail_header_idx + 1, values_only=True):
                    if len(detail_row) <= max(deal_id_idx, repayment_idx):
                        continue
                    deal_id_val = detail_row[deal_id_idx]
                    repayment_val = detail_row[repayment_idx]
                    if deal_id_val and repayment_val:
                        repayment_by_deal_id[str(deal_id_val).strip()] = repayment_val

        keywords = {"borrower", "company", "name", "principal", "instrument", "amount", "rate", "tenure", "status"}
        header_row_values: tuple[Any, ...] | None = None
        rows_seen: list[tuple[Any, ...]] = []

        for row in ws.iter_rows(max_row=20, values_only=True):
            rows_seen.append(row)
            cells_str = [str(c).strip() if c is not None else "" for c in row]
            nonempty = [c for c in cells_str if c]
            if len(nonempty) < 2:
                continue
            row_text = " ".join(cells_str).lower()
            if any(kw in row_text for kw in keywords):
                header_row_values = row
                break
            if len(nonempty) >= 4 and header_row_values is None:
                header_row_values = row

        if header_row_values is None:
            raise HTTPException(status_code=422, detail={"message": "No recognisable header row found.", "missing_columns": []})

        all_columns = [
            str(h).strip() if h is not None else f"Column_{i + 1}"
            for i, h in enumerate(header_row_values)
        ]
        columns = [c for c in all_columns if c]
        missing = _check_schema(columns)
        if missing:
            raise HTTPException(
                status_code=422,
                detail={"message": "Invalid file structure.", "missing_columns": missing, "found_columns": columns},
            )

        header_idx = rows_seen.index(header_row_values)
        _skip_prefixes = ("TOTAL", "SUBTOTAL", "GRAND", "PORTFOLIO", "NOTE", "SOURCE")
        for row in ws.iter_rows(min_row=header_idx + 2, values_only=True):
            nonempty = [v for v in row if v is not None and str(v).strip()]
            if not nonempty:
                continue
            first_cell = str(row[0] if row[0] is not None else "").strip().upper()
            if any(first_cell.startswith(x) for x in _skip_prefixes) or first_cell.startswith("⚠"):
                continue
            parsed_row = {
                col: val
                for col, val in zip(all_columns, row)
            }
            deal_id = str(parsed_row.get("Deal ID") or parsed_row.get("ID") or parsed_row.get("Ref") or "").strip()
            if deal_id and deal_id in repayment_by_deal_id:
                parsed_row["Repayment Structure"] = repayment_by_deal_id[deal_id]
            raw_rows.append(parsed_row)

    if not raw_rows:
        raise HTTPException(status_code=422, detail={"message": "No data rows found.", "missing_columns": []})

    deals = [_norm_row(row, i) for i, row in enumerate(raw_rows)]
    return PortfolioUploadResponse(
        deals=deals,
        filename=filename,
        row_count=len(deals),
    )
